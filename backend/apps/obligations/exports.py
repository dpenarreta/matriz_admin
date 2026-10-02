"""Exportación real a Excel y PDF (secciones 5.4 y 5.13). Cada archivo lleva
la empresa, la fecha de corte y el usuario que exporta, y respeta los
mismos filtros y la misma visibilidad que la pantalla."""

from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.organizations.serializers import full_name

from .status import compute_status

MATRIX_HEADERS = [
    "Código",
    "Obligación",
    "Área",
    "Entidad",
    "Período",
    "Responsable",
    "Vencimiento",
    "Estado",
    "Días",
    "Prioridad",
]


def _days_text(status) -> str:
    if status.group == "incumplido":
        return f"+{status.days_overdue} d atraso"
    if status.group == "finalizado":
        return f"{status.days_overdue} d atraso (cerrado)" if status.is_late else "—"
    if status.is_due_today:
        return "Vence hoy"
    return f"{status.days_remaining} d restantes"


def matrix_rows(periods, now=None):
    now = now or timezone.now()
    for period in periods:
        status = compute_status(period, now)
        tz = ZoneInfo(period.company.timezone)
        yield [
            period.code,
            period.obligation.name,
            period.obligation.area.name,
            period.obligation.control_entity.name,
            period.label,
            full_name(period.responsible),
            period.due_at.astimezone(tz).strftime("%Y-%m-%d %H:%M"),
            status.label,
            _days_text(status),
            period.get_priority_display(),
        ]


def _header_lines(company, user, title):
    cut = timezone.localtime(timezone.now(), ZoneInfo(company.timezone))
    return [
        title,
        f"Empresa: {company.legal_name}",
        f"Fecha de corte: {cut:%Y-%m-%d %H:%M} ({company.timezone})",
        f"Exportado por: {full_name(user)}",
    ]


def to_xlsx(*, company, user, title, headers, rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    for line in _header_lines(company, user, title):
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([])
    sheet.append(headers)
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164B86")
    widths = [len(header) for header in headers]
    for row in rows:
        sheet.append(row)
        widths = [max(width, len(str(value))) for width, value in zip(widths, row, strict=False)]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, 60)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(*, company, user, title, headers, rows) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24
    )
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    lines = _header_lines(company, user, title)
    story = [Paragraph(lines[0], styles["Title"])]
    story += [Paragraph(line, styles["Normal"]) for line in lines[1:]]
    story.append(Spacer(1, 10))
    data = [headers] + [[Paragraph(str(value), cell_style) for value in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#164B86")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C2CBDB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    document.build(story)
    return buffer.getvalue()


def report_rows(report: dict) -> list[list]:
    rows = []
    for section, items in (("Área", report["by_area"]), ("Entidad", report["by_entity"])):
        for item in items:
            rows.append(
                [
                    section,
                    item["name"],
                    item["done"],
                    item["in_progress"],
                    item["overdue"],
                    item["total"],
                ]
            )
    return rows


REPORT_HEADERS = ["Agrupación", "Nombre", "Finalizado", "En progreso", "Incumplido", "Total"]
